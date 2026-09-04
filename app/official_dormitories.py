import base64
import binascii
import hashlib
import io
import re
from typing import Annotated
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, Body, Depends, Query, Request
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session

from .common import (
    all_rows,
    audit,
    authorize,
    authorized_grade_ids,
    clean_text,
    current_user,
    now,
    one,
    require_management,
    require_super_admin,
    require_user,
)
from .database import get_db
from .dormitories import begin_immediate
from .errors import ApiError
from .markdown import markdown_summary, render_markdown
from .rate_limit import enforce_rate_limit

router = APIRouter(prefix="/api/official-dormitories")
admin_router = APIRouter(prefix="/api/admin/official-dormitories")
DB = Annotated[Session, Depends(get_db)]
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSXBody = Annotated[bytes, Body(media_type=XLSX_MEDIA_TYPE)]
MAX_XLSX_BYTES = 2 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_XLSX_ENTRIES = 1000
IMPORT_HEADERS = (
    "宿舍编号",
    "成员1登录标识（宿舍长）",
    "成员2登录标识",
    "成员3登录标识",
    "成员4登录标识",
)
DORMITORY_CODE = re.compile(r"[\w\-\u4e00-\u9fff]{1,40}", re.ASCII)
CONTENT_FIELDS = {
    "nickname": ("NICKNAME", "nickname", "nickname_status", 20),
    "description": ("DESCRIPTION", "description_markdown", "description_status", 2000),
    "rules": ("RULES", "rules_markdown", "rules_status", 5000),
}
OFFICIAL_NOT_FOUND_MESSAGE = "宿舍不存在或暂不可访问"
ADMIN_NOT_FOUND_MESSAGE = "正式宿舍不存在"


def active_student(user: dict) -> None:
    require_user(user)
    if user["status"] != "ACTIVE":
        raise ApiError(403, "ACCOUNT_UNAVAILABLE", "账号当前不可使用宿得道理")


def enforce_read_rate(request: Request, user_id: int) -> None:
    ip_address = request.client.host if request.client else "unknown"
    enforce_rate_limit("official-dormitory-read-user", str(user_id), 120, 60, "DORMITORY_READ_RATE_LIMITED")
    enforce_rate_limit("official-dormitory-read-ip", ip_address, 600, 60, "DORMITORY_READ_RATE_LIMITED")


def enforce_write_rate(request: Request, user_id: int, operation: str, count: int, ip_count: int, period: int) -> None:
    ip_address = request.client.host if request.client else "unknown"
    enforce_rate_limit(f"official-dormitory-{operation}-user", str(user_id), count, period, "DORMITORY_RATE_LIMITED")
    enforce_rate_limit(f"official-dormitory-{operation}-ip", ip_address, ip_count, period, "DORMITORY_RATE_LIMITED")


def encoded_cursor(row: dict) -> str:
    raw = f"{int(row['is_mine'])}|{row['updated_at']}|{row['id']}".encode()
    return base64.urlsafe_b64encode(raw).decode().rstrip("=")


def decoded_cursor(cursor: str) -> tuple[int, str, int] | None:
    if not cursor:
        return None
    try:
        raw = base64.urlsafe_b64decode(cursor + "=" * (-len(cursor) % 4)).decode()
        mine, updated_at, dormitory_id = raw.split("|", 2)
        if mine not in ("0", "1") or not updated_at or int(dormitory_id) < 1:
            raise ValueError
        return int(mine), updated_at, int(dormitory_id)
    except (binascii.Error, ValueError):
        raise ApiError(400, "INVALID_DORMITORY_CURSOR", "分页位置无效") from None


def dormitory_by_id(db: Session, dormitory_id: int) -> dict | None:
    return one(db, "SELECT * FROM official_dormitories WHERE id=:id", {"id": dormitory_id})


def member_for_user(db: Session, dormitory_id: int, user_id: int) -> dict | None:
    return one(
        db,
        """SELECT * FROM official_dormitory_members
        WHERE official_dormitory_id=:dormitory AND user_id=:user""",
        {"dormitory": dormitory_id, "user": user_id},
    )


def member_rows(db: Session, dormitory_ids: list[int]) -> list[dict]:
    if not dormitory_ids:
        return []
    parameters = {f"id_{index}": value for index, value in enumerate(dormitory_ids)}
    placeholders = ",".join(f":{key}" for key in parameters)
    return all_rows(
        db,
        """SELECT member.id,member.official_dormitory_id,member.user_id,member.role,member.position,
          member.imported_login_identifier,member.name_snapshot,member.grade_snapshot,member.major_snapshot,
          user.login_identifier,user.name,user.grade,user.major,user.gender,user.status AS user_status,
          card.id AS card_id,card.status AS card_status,card.avatar_url,card.origin_city,card.one_sentence_intro
        FROM official_dormitory_members member
        LEFT JOIN users user ON user.id=member.user_id
        LEFT JOIN roommate_cards card ON card.user_id=member.user_id
        WHERE member.official_dormitory_id IN ("""
        + placeholders
        + ") ORDER BY member.official_dormitory_id,member.position",
        parameters,
    )


def blocked_user_ids(db: Session, viewer_id: int, user_ids: list[int]) -> set[int]:
    if not user_ids:
        return set()
    parameters = {f"user_{index}": value for index, value in enumerate(dict.fromkeys(user_ids))}
    placeholders = ",".join(f":{key}" for key in parameters)
    rows = all_rows(
        db,
        f"""SELECT CASE WHEN blocker_id=:viewer THEN blocked_id ELSE blocker_id END AS user_id
        FROM blocks WHERE (blocker_id=:viewer AND blocked_id IN ({placeholders}))
        OR (blocked_id=:viewer AND blocker_id IN ({placeholders}))""",
        {**parameters, "viewer": viewer_id},
    )
    return {row["user_id"] for row in rows}


def student_dormitories(db: Session, rows: list[dict], viewer_id: int, details: bool = False) -> list[dict]:
    members = member_rows(db, [row["id"] for row in rows])
    blocked = blocked_user_ids(db, viewer_id, [row["user_id"] for row in members if row["user_id"]])
    grouped = {row["id"]: [] for row in rows}
    own_ids = {row["id"] for row in rows if row.get("is_mine")}
    for member in members:
        own_dormitory = member["official_dormitory_id"] in own_ids
        visible = own_dormitory or bool(
            member["user_id"]
            and member["user_status"] == "ACTIVE"
            and member["card_status"] == "PUBLISHED"
            and member["user_id"] not in blocked
        )
        item = {
            "position": member["position"],
            "role": member["role"],
            "visible": visible,
            "leaderPending": member["role"] == "LEADER" and not member["user_id"],
        }
        if visible:
            item.update(
                {
                    "memberId": member["id"] if own_dormitory else None,
                    "cardId": member["card_id"] if member["card_status"] == "PUBLISHED" else None,
                    "cardStatus": member["card_status"] or "UNPUBLISHED",
                    "isOwnCard": member["user_id"] == viewer_id,
                    "name": member["name"] or member["name_snapshot"] or "已删除账号",
                    "grade": member["grade"] or member["grade_snapshot"],
                    "major": member["major"] if member["user_id"] else member["major_snapshot"],
                    "gender": member["gender"] or "UNSPECIFIED",
                    "avatarUrl": member["avatar_url"] if member["card_status"] == "PUBLISHED" else "",
                    "originCity": member["origin_city"] if member["card_status"] == "PUBLISHED" else "",
                    "introduction": member["one_sentence_intro"] if member["card_status"] == "PUBLISHED" else "",
                }
            )
        grouped[member["official_dormitory_id"]].append(item)
    result = []
    for row in rows:
        own = bool(row.get("is_mine"))
        nickname = row["nickname"] if row["nickname_status"] == "NORMAL" else ""
        description = row["description_markdown"] if row["description_status"] == "NORMAL" else ""
        summary, truncated = markdown_summary(description, 120) if description else ("", False)
        item = {
            "id": row["id"],
            "dormitoryCode": row["dormitory_code"],
            "nickname": nickname,
            "displayName": f"{row['dormitory_code']}·{nickname}" if nickname else row["dormitory_code"],
            "descriptionSummary": summary,
            "descriptionSummaryTruncated": truncated,
            "memberCount": row.get("member_count", len(grouped[row["id"]])),
            "isMine": own,
            "version": row["version"],
            "updatedAt": row["updated_at"],
            "members": grouped[row["id"]],
        }
        if details:
            rules = row["rules_markdown"] if row["rules_status"] == "NORMAL" else ""
            rules_summary, rules_truncated = markdown_summary(rules, 120) if rules else ("", False)
            item.update(
                {
                    "nicknameHidden": row["nickname_status"] != "NORMAL",
                    "descriptionHidden": row["description_status"] != "NORMAL",
                    "rulesHidden": row["rules_status"] != "NORMAL",
                    "descriptionHtml": render_markdown(description) if description else "",
                    "rulesHtml": render_markdown(rules) if rules else "",
                    "rulesSummary": rules_summary,
                    "rulesSummaryTruncated": rules_truncated,
                    "descriptionMarkdown": description if own else None,
                    "rulesMarkdown": rules if own else None,
                    "canEdit": own,
                    "canTransferLeader": any(
                        member.get("memberId") and member.get("isOwnCard") and member["role"] == "LEADER"
                        for member in grouped[row["id"]]
                    ),
                }
            )
        result.append(item)
    return result


@router.get("/mine")
def my_official_dormitories(request: Request, db: DB) -> dict:
    user = current_user(request, db)
    active_student(user)
    enforce_read_rate(request, user["id"])
    rows = all_rows(
        db,
        """SELECT dormitory.*,
          1 AS is_mine,(SELECT COUNT(*) FROM official_dormitory_members
          WHERE official_dormitory_id=dormitory.id) AS member_count
        FROM official_dormitories dormitory JOIN official_dormitory_members member
          ON member.official_dormitory_id=dormitory.id
        WHERE member.user_id=:user ORDER BY member.created_at DESC,dormitory.id""",
        {"user": user["id"]},
    )
    return {"dormitories": student_dormitories(db, rows, user["id"])}


@router.get("")
def list_official_dormitories(
    request: Request,
    db: DB,
    search: str = "",
    cursor: str = "",
    limit: Annotated[int, Query(ge=1, le=30)] = 15,
) -> dict:
    user = current_user(request, db)
    active_student(user)
    enforce_read_rate(request, user["id"])
    query = clean_text(search, 80)
    own = "EXISTS(SELECT 1 FROM official_dormitory_members own WHERE own.official_dormitory_id=dormitory.id AND own.user_id=:viewer)"
    params: dict = {"viewer": user["id"], "limit": limit + 1}
    filters = []
    if query:
        params["search"] = f"%{query}%"
        filters.append(
            f"""(LOWER(dormitory.dormitory_code) LIKE LOWER(:search)
            OR LOWER(dormitory.nickname) LIKE LOWER(:search)
            OR EXISTS(SELECT 1 FROM official_dormitory_members searched
              LEFT JOIN users searched_user ON searched_user.id=searched.user_id
              LEFT JOIN roommate_cards searched_card ON searched_card.user_id=searched.user_id
              WHERE searched.official_dormitory_id=dormitory.id
              AND LOWER(COALESCE(searched_user.name,searched.name_snapshot)) LIKE LOWER(:search)
              AND ({own} OR (searched_user.status='ACTIVE' AND searched_card.status='PUBLISHED'
                AND NOT EXISTS(SELECT 1 FROM blocks block WHERE
                  (block.blocker_id=:viewer AND block.blocked_id=searched.user_id)
                  OR (block.blocker_id=searched.user_id AND block.blocked_id=:viewer))))))"""
        )
    parsed_cursor = decoded_cursor(cursor)
    if parsed_cursor:
        params.update({"cursor_mine": parsed_cursor[0], "cursor_time": parsed_cursor[1], "cursor_id": parsed_cursor[2]})
        filters.append(
            f"""(({own})<:cursor_mine OR (({own})=:cursor_mine AND
            (dormitory.updated_at<:cursor_time OR
            (dormitory.updated_at=:cursor_time AND dormitory.id<:cursor_id))))"""
        )
    where = " AND ".join(filters) if filters else "1=1"
    rows = all_rows(
        db,
        f"""SELECT dormitory.*,({own}) AS is_mine,
          (SELECT COUNT(*) FROM official_dormitory_members WHERE official_dormitory_id=dormitory.id) AS member_count
        FROM official_dormitories dormitory WHERE {where}
        ORDER BY is_mine DESC,dormitory.updated_at DESC,dormitory.id DESC LIMIT :limit""",
        params,
    )
    has_more = len(rows) > limit
    rows = rows[:limit]
    return {
        "dormitories": student_dormitories(db, rows, user["id"]),
        "nextCursor": encoded_cursor(rows[-1]) if has_more and rows else None,
    }


@router.get("/{dormitory_id}")
def official_dormitory_detail(dormitory_id: int, request: Request, db: DB) -> dict:
    user = current_user(request, db)
    active_student(user)
    enforce_read_rate(request, user["id"])
    row = one(
        db,
        """SELECT dormitory.*,
          EXISTS(SELECT 1 FROM official_dormitory_members member
            WHERE member.official_dormitory_id=dormitory.id AND member.user_id=:viewer) AS is_mine,
          (SELECT COUNT(*) FROM official_dormitory_members
            WHERE official_dormitory_id=dormitory.id) AS member_count
        FROM official_dormitories dormitory WHERE dormitory.id=:id""",
        {"id": dormitory_id, "viewer": user["id"]},
    )
    if not row:
        raise ApiError(404, "OFFICIAL_DORMITORY_NOT_FOUND", OFFICIAL_NOT_FOUND_MESSAGE)
    return {"dormitory": student_dormitories(db, [row], user["id"], True)[0]}


def validated_content(body: dict, row: dict) -> list[tuple[str, str, str, str]]:
    changes = []
    for request_field, (revision_field, column, status_column, maximum) in CONTENT_FIELDS.items():
        if request_field not in body:
            continue
        if row[status_column] != "NORMAL":
            raise ApiError(409, "DORMITORY_CONTENT_HIDDEN", "被治理的内容不能直接修改")
        value = clean_text(body.get(request_field), maximum)
        if request_field == "nickname" and ("·" in value or any(ord(char) < 32 for char in value)):
            raise ApiError(400, "INVALID_DORMITORY_NICKNAME", "宿舍昵称包含不受支持的字符")
        if value != row[column]:
            changes.append((revision_field, column, row[column], value))
    if not changes and not any(field in body for field in CONTENT_FIELDS):
        raise ApiError(400, "DORMITORY_CONTENT_REQUIRED", "请选择需要修改的内容")
    return changes


@router.patch("/{dormitory_id}")
def update_official_dormitory(dormitory_id: int, request: Request, body: dict, db: DB) -> dict:
    user = current_user(request, db)
    active_student(user)
    enforce_write_rate(request, user["id"], "content", 20, 100, 60)
    if type(body.get("version")) is not int:
        raise ApiError(400, "DORMITORY_VERSION_REQUIRED", "请提供宿舍内容版本")
    begin_immediate(db)
    row = dormitory_by_id(db, dormitory_id)
    if not row or not member_for_user(db, dormitory_id, user["id"]):
        raise ApiError(404, "OFFICIAL_DORMITORY_NOT_FOUND", OFFICIAL_NOT_FOUND_MESSAGE)
    if row["version"] != body["version"]:
        raise ApiError(409, "DORMITORY_CONTENT_CONFLICT", "宿舍内容已被其他成员修改，请重新载入")
    changes = validated_content(body, row)
    if changes:
        timestamp = now()
        next_version = row["version"] + 1
        assignments = []
        params = {"id": dormitory_id, "version": next_version, "now": timestamp}
        for index, (field_name, column, previous, value) in enumerate(changes):
            key = f"value_{index}"
            assignments.append(f"{column}=:{key}")
            params[key] = value
            db.execute(
                text(
                    """INSERT INTO official_dormitory_revisions(
                    official_dormitory_id,field_name,previous_value,new_value,from_version,to_version,
                    edited_by,editor_name_snapshot,created_at)
                    VALUES(:dormitory,:field,:previous,:new,:from_version,:to_version,:editor,:name,:now)"""
                ),
                {
                    "dormitory": dormitory_id,
                    "field": field_name,
                    "previous": previous,
                    "new": value,
                    "from_version": row["version"],
                    "to_version": next_version,
                    "editor": user["id"],
                    "name": user["name"],
                    "now": timestamp,
                },
            )
        db.execute(
            text(
                f"UPDATE official_dormitories SET {','.join(assignments)},version=:version,updated_at=:now WHERE id=:id"
            ),
            params,
        )
        audit(
            db,
            user,
            request,
            "UPDATE_OFFICIAL_DORMITORY_CONTENT",
            "OFFICIAL_DORMITORY",
            dormitory_id,
            metadata={"fields": [item[0] for item in changes], "version": next_version},
        )
    db.commit()
    return {"ok": True, "version": row["version"] + (1 if changes else 0)}


@router.post("/markdown/preview")
def preview_official_dormitory_markdown(request: Request, body: dict, db: DB) -> dict:
    user = current_user(request, db)
    active_student(user)
    enforce_write_rate(request, user["id"], "preview", 30, 120, 60)
    field = body.get("field")
    maximum = 0
    if field == "description":
        maximum = 2000
    elif field == "rules":
        maximum = 5000
    if not maximum:
        raise ApiError(400, "INVALID_DORMITORY_CONTENT_FIELD", "预览字段无效")
    content = clean_text(body.get("content"), maximum)
    return {"html": render_markdown(content) if content else ""}


@router.post("/{dormitory_id}/leader-transfer")
def transfer_official_dormitory_leader(dormitory_id: int, request: Request, body: dict, db: DB) -> dict:
    user = current_user(request, db)
    active_student(user)
    enforce_write_rate(request, user["id"], "leader", 5, 30, 600)
    target_id = body.get("targetMemberId")
    version = body.get("version")
    if type(target_id) is not int or type(version) is not int:
        raise ApiError(400, "INVALID_LEADER_TRANSFER", "宿舍长转让信息无效")
    begin_immediate(db)
    dormitory = dormitory_by_id(db, dormitory_id)
    current = member_for_user(db, dormitory_id, user["id"])
    target = one(
        db,
        """SELECT member.*,user.status,user.account_type,user.name FROM official_dormitory_members member
        JOIN users user ON user.id=member.user_id
        WHERE member.id=:member AND member.official_dormitory_id=:dormitory""",
        {"member": target_id, "dormitory": dormitory_id},
    )
    if not dormitory or not current:
        raise ApiError(404, "OFFICIAL_DORMITORY_NOT_FOUND", OFFICIAL_NOT_FOUND_MESSAGE)
    if dormitory["version"] != version:
        raise ApiError(409, "DORMITORY_CONTENT_CONFLICT", "宿舍信息已变化，请重新载入")
    if current["role"] != "LEADER":
        raise ApiError(403, "DORMITORY_LEADER_REQUIRED", "只有宿舍长可以转让身份")
    if (
        not target
        or target["user_id"] == user["id"]
        or target["status"] != "ACTIVE"
        or target["account_type"] != "USER"
    ):
        raise ApiError(400, "INVALID_LEADER_TARGET", "只能转让给同宿舍的其他有效成员")
    timestamp = now()
    db.execute(text("UPDATE official_dormitory_members SET role='MEMBER' WHERE id=:id"), {"id": current["id"]})
    db.execute(text("UPDATE official_dormitory_members SET role='LEADER' WHERE id=:id"), {"id": target["id"]})
    db.execute(
        text("UPDATE official_dormitories SET version=version+1,updated_at=:now WHERE id=:id"),
        {"id": dormitory_id, "now": timestamp},
    )
    audit(
        db,
        user,
        request,
        "TRANSFER_OFFICIAL_DORMITORY_LEADER",
        "OFFICIAL_DORMITORY",
        dormitory_id,
        metadata={"fromMemberId": current["id"], "toMemberId": target["id"]},
    )
    db.commit()
    return {"ok": True, "version": dormitory["version"] + 1}


def management_user(request: Request, db: Session) -> dict:
    user = current_user(request, db)
    require_management(db, user)
    return user


def spreadsheet_value(cell: object, row_number: int, column_number: int) -> str:
    data_type = getattr(cell, "data_type", "")
    value = getattr(cell, "value", None)
    if data_type in ("e", "f") or getattr(cell, "hyperlink", None):
        raise ApiError(400, "INVALID_DORMITORY_IMPORT", f"第 {row_number} 行第 {column_number} 列包含不安全内容")
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ApiError(400, "INVALID_DORMITORY_IMPORT", f"第 {row_number} 行第 {column_number} 列必须设置为文本")
    return value.strip()


def validate_xlsx_package(content: bytes) -> None:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            total_size = sum(entry.file_size for entry in entries)
            forbidden = (
                "xl/activex/",
                "xl/drawings/",
                "xl/embeddings/",
                "xl/externallinks/",
                "xl/media/",
                "xl/oleobjects/",
            )
            if len(entries) > MAX_XLSX_ENTRIES or total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ApiError(400, "INVALID_DORMITORY_IMPORT_FILE", "Excel 文件结构或解压体积超出限制")
            unsafe = any(
                entry.flag_bits & 1
                or entry.filename.replace("\\", "/").lower().endswith("vbaproject.bin")
                or entry.filename.replace("\\", "/").lower().startswith(forbidden)
                for entry in entries
            )
            external = any(
                entry.filename.lower().endswith(".rels")
                and re.search(rb"targetmode\s*=\s*(['\"])external\1", archive.read(entry), re.IGNORECASE)
                for entry in entries
            )
            unsafe_xml = any(
                entry.filename.lower().endswith((".xml", ".rels"))
                and re.search(rb"<!\s*(doctype|entity)", archive.read(entry), re.IGNORECASE)
                for entry in entries
            )
            if unsafe or external or unsafe_xml:
                raise ApiError(400, "INVALID_DORMITORY_IMPORT_FILE", "Excel 文件包含不支持的外部或嵌入内容")
    except BadZipFile as error:
        raise ApiError(400, "INVALID_DORMITORY_IMPORT_FILE", "无法读取该 Excel 文件") from error


def xlsx_rows(content: bytes, filename: str) -> list[dict]:
    if not filename.lower().endswith(".xlsx"):
        raise ApiError(400, "INVALID_DORMITORY_IMPORT_FILE", "请选择 .xlsx 文件")
    if not content or len(content) > MAX_XLSX_BYTES:
        raise ApiError(413, "DORMITORY_IMPORT_TOO_LARGE", "导入文件不能超过 2 MiB")
    validate_xlsx_package(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as error:
        raise ApiError(400, "INVALID_DORMITORY_IMPORT_FILE", "无法读取该 Excel 文件") from error
    try:
        if not workbook.worksheets:
            raise ApiError(400, "INVALID_DORMITORY_IMPORT_FILE", "Excel 文件没有工作表")
        sheet = workbook.worksheets[0]
        if sheet.max_column > len(IMPORT_HEADERS):
            raise ApiError(400, "INVALID_DORMITORY_IMPORT", "导入表包含模板以外的列")
        rows = sheet.iter_rows(min_col=1, max_col=len(IMPORT_HEADERS))
        header = tuple(spreadsheet_value(cell, 1, index) for index, cell in enumerate(next(rows, ()), 1))
        if header != IMPORT_HEADERS:
            raise ApiError(400, "INVALID_DORMITORY_IMPORT", "导入表头与模板不一致")
        result = []
        for row_number, cells in enumerate(rows, 2):
            if row_number > 2001:
                raise ApiError(400, "DORMITORY_IMPORT_TOO_MANY_ROWS", "导入数据不能超过 2000 行")
            values = [spreadsheet_value(cell, row_number, index) for index, cell in enumerate(cells, 1)]
            if not any(values):
                continue
            code, *identifiers = values
            if not code or not identifiers[0]:
                raise ApiError(400, "INVALID_DORMITORY_IMPORT", f"第 {row_number} 行缺少宿舍编号或宿舍长")
            if not DORMITORY_CODE.fullmatch(code) or code.startswith(("=", "+", "-", "@")):
                raise ApiError(400, "INVALID_DORMITORY_CODE", f"第 {row_number} 行宿舍编号格式无效")
            present = [identifier for identifier in identifiers if identifier]
            if identifiers[: len(present)] != present or len(present) != len(set(present)):
                raise ApiError(400, "INVALID_DORMITORY_IMPORT", f"第 {row_number} 行成员存在跳列或重复")
            result.append({"rowNumber": row_number, "dormitoryCode": code, "identifiers": present})
        if not result:
            raise ApiError(400, "DORMITORY_IMPORT_EMPTY", "导入表中没有宿舍数据")
        codes = [row["dormitoryCode"] for row in result]
        if len(codes) != len(set(codes)):
            raise ApiError(400, "DUPLICATE_DORMITORY_CODE", "导入文件包含重复宿舍编号")
        return result
    finally:
        workbook.close()


def import_accounts(db: Session, rows: list[dict]) -> dict[str, dict]:
    identifiers = list(dict.fromkeys(identifier for row in rows for identifier in row["identifiers"]))
    accounts = []
    for start in range(0, len(identifiers), 400):
        chunk = identifiers[start : start + 400]
        params = {f"login_{index}": value for index, value in enumerate(chunk)}
        accounts.extend(
            all_rows(
                db,
                """SELECT id,login_identifier,name,grade,grade_id,major,status,account_type FROM users
                WHERE login_identifier IN ("""
                + ",".join(f":{key}" for key in params)
                + ")",
                params,
            )
        )
    return {account["login_identifier"]: account for account in accounts}


def analyze_import(db: Session, admin: dict, rows: list[dict]) -> list[dict]:
    accounts = import_accounts(db, rows)
    analyzed = []
    for row in rows:
        members = [accounts.get(identifier) for identifier in row["identifiers"]]
        if any(member is None for member in members):
            missing = [
                identifier for identifier, member in zip(row["identifiers"], members, strict=True) if member is None
            ]
            raise ApiError(
                400, "DORMITORY_IMPORT_USER_NOT_FOUND", f"第 {row['rowNumber']} 行账号不存在：{'、'.join(missing)}"
            )
        typed_members = [member for member in members if member]
        if any(
            member["account_type"] != "USER" or member["status"] not in ("PENDING_ACTIVATION", "ACTIVE")
            for member in typed_members
        ):
            raise ApiError(400, "DORMITORY_IMPORT_USER_INVALID", f"第 {row['rowNumber']} 行包含不可导入账号")
        grade_ids = {member["grade_id"] for member in typed_members}
        if None in grade_ids or len(grade_ids) != 1:
            raise ApiError(400, "DORMITORY_IMPORT_GRADE_MISMATCH", f"第 {row['rowNumber']} 行成员必须属于同一年级")
        grade_id = next(iter(grade_ids))
        grant = authorize(db, admin, "OFFICIAL_DORMITORY_IMPORT", grade_id)
        analyzed.append({**row, "members": typed_members, "managementGradeId": grade_id, "grant": grant})
    return analyzed


def existing_import_state(db: Session, analyzed: list[dict]) -> dict[str, dict]:
    codes = [row["dormitoryCode"] for row in analyzed]
    existing = []
    for start in range(0, len(codes), 400):
        chunk = codes[start : start + 400]
        params = {f"code_{index}": value for index, value in enumerate(chunk)}
        existing.extend(
            all_rows(
                db,
                "SELECT * FROM official_dormitories WHERE dormitory_code IN ("
                + ",".join(f":{key}" for key in params)
                + ")",
                params,
            )
        )
    by_code = {row["dormitory_code"]: row for row in existing}
    members = member_rows(db, [row["id"] for row in existing])
    for row in existing:
        row["member_user_ids"] = [
            member["user_id"] for member in members if member["official_dormitory_id"] == row["id"]
        ]
    return by_code


def import_plan(db: Session, analyzed: list[dict]) -> tuple[list[dict], list[dict]]:
    existing = existing_import_state(db, analyzed)
    create = []
    skip = []
    for row in analyzed:
        current = existing.get(row["dormitoryCode"])
        if not current:
            create.append(row)
            continue
        desired = [member["id"] for member in row["members"]]
        if current["management_grade_id"] == row["managementGradeId"] and current["member_user_ids"] == desired:
            skip.append(row)
            continue
        raise ApiError(409, "OFFICIAL_DORMITORY_EXISTS", f"宿舍 {row['dormitoryCode']} 已存在且成员不同")
    return create, skip


def import_request(
    request: Request, content: bytes, db: Session, admin: dict
) -> tuple[bytes, list[dict], list[dict], list[dict]]:
    if authorized_grade_ids(db, admin, "OFFICIAL_DORMITORY_IMPORT") == []:
        raise ApiError(403, "PERMISSION_DENIED", "当前账号缺少所需管理权限")
    content_type = request.headers.get("content-type", "").partition(";")[0].strip().lower()
    if content_type != XLSX_MEDIA_TYPE:
        raise ApiError(415, "INVALID_DORMITORY_IMPORT_FILE", "请上传 .xlsx 格式的 Excel 文件")
    rows = xlsx_rows(content, request.headers.get("x-file-name", ""))
    analyzed = analyze_import(db, admin, rows)
    create, skip = import_plan(db, analyzed)
    return content, analyzed, create, skip


@admin_router.post("/import/preview")
def preview_import(request: Request, content: XLSXBody, db: DB) -> dict:
    admin = management_user(request, db)
    enforce_write_rate(request, admin["id"], "import-preview", 10, 30, 600)
    content, analyzed, create, skip = import_request(request, content, db, admin)
    create_codes = {row["dormitoryCode"] for row in create}
    return {
        "sha256": hashlib.sha256(content).hexdigest(),
        "rowCount": len(analyzed),
        "createCount": len(create),
        "skipCount": len(skip),
        "rows": [
            {
                "rowNumber": row["rowNumber"],
                "dormitoryCode": row["dormitoryCode"],
                "grade": row["members"][0]["grade"],
                "members": [member["name"] for member in row["members"]],
                "action": "CREATE" if row["dormitoryCode"] in create_codes else "SKIP",
            }
            for row in analyzed
        ],
    }


@admin_router.post("/import/confirm")
def confirm_import(request: Request, content: XLSXBody, db: DB) -> dict:
    admin = management_user(request, db)
    enforce_write_rate(request, admin["id"], "import-confirm", 5, 20, 600)
    expected = request.headers.get("x-file-sha256", "")
    actual = hashlib.sha256(content).hexdigest()
    if not expected or expected != actual:
        raise ApiError(409, "DORMITORY_IMPORT_FILE_CHANGED", "文件与预检时不一致，请重新预检")
    rows = xlsx_rows(content, request.headers.get("x-file-name", ""))
    try:
        begin_immediate(db)
        analyzed = analyze_import(db, admin, rows)
        create, skip = import_plan(db, analyzed)
        timestamp = now()
        for row in create:
            result = db.execute(
                text(
                    """INSERT INTO official_dormitories(
                    dormitory_code,management_grade_id,created_by,created_at,updated_at)
                    VALUES(:code,:grade,:admin,:now,:now)"""
                ),
                {
                    "code": row["dormitoryCode"],
                    "grade": row["managementGradeId"],
                    "admin": admin["id"],
                    "now": timestamp,
                },
            )
            dormitory_id = result.lastrowid
            for position, member in enumerate(row["members"], 1):
                db.execute(
                    text(
                        """INSERT INTO official_dormitory_members(
                        official_dormitory_id,user_id,role,position,imported_login_identifier,
                        name_snapshot,grade_snapshot,major_snapshot,created_at)
                        VALUES(:dormitory,:user,:role,:position,:login,:name,:grade,:major,:now)"""
                    ),
                    {
                        "dormitory": dormitory_id,
                        "user": member["id"],
                        "role": "LEADER" if position == 1 else "MEMBER",
                        "position": position,
                        "login": member["login_identifier"],
                        "name": member["name"],
                        "grade": member["grade"],
                        "major": member["major"],
                        "now": timestamp,
                    },
                )
            audit(
                db,
                admin,
                request,
                "IMPORT_OFFICIAL_DORMITORY",
                "OFFICIAL_DORMITORY",
                dormitory_id,
                metadata={"fileSha256": actual, "memberCount": len(row["members"])},
                grant=row["grant"],
                after={"dormitoryCode": row["dormitoryCode"]},
            )
        db.commit()
    except (IntegrityError, OperationalError) as error:
        db.rollback()
        raise ApiError(409, "DORMITORY_IMPORT_CONFLICT", "正式宿舍名单已变化，请重新预检") from error
    return {"created": len(create), "skipped": len(skip)}


def admin_scope_sql(grade_ids: list[int] | None, params: dict) -> str:
    if grade_ids is None:
        return "1=1"
    if not grade_ids:
        raise ApiError(403, "PERMISSION_DENIED", "当前账号缺少所需管理权限")
    keys = []
    for index, grade_id in enumerate(grade_ids):
        key = f"grade_{index}"
        params[key] = grade_id
        keys.append(f":{key}")
    return f"dormitory.management_grade_id IN ({','.join(keys)})"


@admin_router.get("")
def admin_dormitories(
    request: Request,
    db: DB,
    search: str = "",
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=50)] = 30,
) -> dict:
    admin = management_user(request, db)
    enforce_read_rate(request, admin["id"])
    query = clean_text(search, 80)
    params: dict = {"offset": offset, "limit": limit}
    scope = admin_scope_sql(authorized_grade_ids(db, admin, "OFFICIAL_DORMITORY_READ"), params)
    search_sql = ""
    if query:
        params["search"] = f"%{query}%"
        search_sql = """AND (LOWER(dormitory.dormitory_code) LIKE LOWER(:search)
          OR LOWER(dormitory.nickname) LIKE LOWER(:search)
          OR EXISTS(SELECT 1 FROM official_dormitory_members member LEFT JOIN users user ON user.id=member.user_id
            WHERE member.official_dormitory_id=dormitory.id
            AND LOWER(COALESCE(user.name,member.name_snapshot)) LIKE LOWER(:search)))"""
    rows = all_rows(
        db,
        f"""SELECT dormitory.*,grade.name AS management_grade,
          (SELECT COUNT(*) FROM official_dormitory_members WHERE official_dormitory_id=dormitory.id) AS member_count
        FROM official_dormitories dormitory JOIN grades grade ON grade.id=dormitory.management_grade_id
        WHERE {scope} {search_sql} ORDER BY dormitory.id DESC LIMIT :limit OFFSET :offset""",
        params,
    )
    count = one(
        db,
        f"""SELECT COUNT(*) AS total FROM official_dormitories dormitory
        WHERE {scope} {search_sql}""",
        params,
    )
    members = member_rows(db, [row["id"] for row in rows])
    by_id = {row["id"]: [] for row in rows}
    for member in members:
        by_id[member["official_dormitory_id"]].append(
            {
                "memberId": member["id"],
                "name": member["name"] or member["name_snapshot"],
                "grade": member["grade"] or member["grade_snapshot"],
                "role": member["role"],
                "leaderPending": member["role"] == "LEADER" and not member["user_id"],
            }
        )
    for row in rows:
        row["members"] = by_id[row["id"]]
    return {"dormitories": rows, "total": count["total"] if count else 0}


def admin_dormitory(db: Session, admin: dict, dormitory_id: int, permission: str) -> tuple[dict, dict]:
    dormitory = dormitory_by_id(db, dormitory_id)
    if not dormitory:
        raise ApiError(404, "OFFICIAL_DORMITORY_NOT_FOUND", ADMIN_NOT_FOUND_MESSAGE)
    grant = authorize(db, admin, permission, dormitory["management_grade_id"])
    return dormitory, grant


@admin_router.get("/{dormitory_id}")
def admin_dormitory_detail(dormitory_id: int, request: Request, db: DB) -> dict:
    admin = management_user(request, db)
    enforce_read_rate(request, admin["id"])
    dormitory, _ = admin_dormitory(db, admin, dormitory_id, "OFFICIAL_DORMITORY_READ")
    members = member_rows(db, [dormitory_id])
    revisions = all_rows(
        db,
        """SELECT id,field_name,from_version,to_version,editor_name_snapshot,created_at
        FROM official_dormitory_revisions WHERE official_dormitory_id=:id ORDER BY id DESC LIMIT 100""",
        {"id": dormitory_id},
    )
    return {
        "dormitory": {
            **dormitory,
            "descriptionHtml": render_markdown(dormitory["description_markdown"]),
            "rulesHtml": render_markdown(dormitory["rules_markdown"]),
            "members": [
                {
                    "memberId": member["id"],
                    "loginIdentifier": member["login_identifier"] or member["imported_login_identifier"],
                    "importedLoginIdentifier": member["imported_login_identifier"],
                    "name": member["name"] or member["name_snapshot"],
                    "grade": member["grade"] or member["grade_snapshot"],
                    "major": member["major"] if member["user_id"] else member["major_snapshot"],
                    "role": member["role"],
                    "position": member["position"],
                    "leaderPending": member["role"] == "LEADER" and not member["user_id"],
                }
                for member in members
            ],
            "revisions": revisions,
        }
    }


def validated_member_accounts(db: Session, identifiers: object) -> list[dict]:
    if not isinstance(identifiers, list) or not 1 <= len(identifiers) <= 4:
        raise ApiError(400, "INVALID_OFFICIAL_DORMITORY_MEMBERS", "正式宿舍需要 1 至 4 名成员")
    cleaned = [clean_text(identifier, 100, True) for identifier in identifiers]
    if len(cleaned) != len(set(cleaned)):
        raise ApiError(400, "INVALID_OFFICIAL_DORMITORY_MEMBERS", "正式宿舍成员不能重复")
    rows = [{"identifiers": cleaned}]
    accounts = import_accounts(db, rows)
    members = [accounts.get(identifier) for identifier in cleaned]
    if any(member is None for member in members):
        raise ApiError(400, "DORMITORY_IMPORT_USER_NOT_FOUND", "正式宿舍成员账号不存在")
    result = [member for member in members if member]
    if any(
        member["account_type"] != "USER" or member["status"] not in ("PENDING_ACTIVATION", "ACTIVE")
        for member in result
    ):
        raise ApiError(400, "DORMITORY_IMPORT_USER_INVALID", "正式宿舍包含不可用账号")
    if len({member["grade_id"] for member in result}) != 1 or result[0]["grade_id"] is None:
        raise ApiError(400, "DORMITORY_IMPORT_GRADE_MISMATCH", "正式宿舍成员必须属于同一年级")
    return result


@admin_router.put("/{dormitory_id}/members")
def replace_official_dormitory_members(dormitory_id: int, request: Request, body: dict, db: DB) -> dict:
    admin = management_user(request, db)
    enforce_write_rate(request, admin["id"], "members", 5, 20, 600)
    reason = clean_text(body.get("reason"), 200, True)
    version = body.get("version")
    if type(version) is not int:
        raise ApiError(400, "DORMITORY_VERSION_REQUIRED", "请提供宿舍信息版本")
    begin_immediate(db)
    dormitory = dormitory_by_id(db, dormitory_id)
    if not dormitory:
        raise ApiError(404, "OFFICIAL_DORMITORY_NOT_FOUND", ADMIN_NOT_FOUND_MESSAGE)
    authorize(db, admin, "OFFICIAL_DORMITORY_MEMBER_UPDATE", dormitory["management_grade_id"])
    if body.get("confirmation") != dormitory["dormitory_code"]:
        raise ApiError(400, "CONFIRMATION_REQUIRED", "请输入宿舍编号确认成员纠正")
    members = validated_member_accounts(db, body.get("loginIdentifiers"))
    new_grade = members[0]["grade_id"]
    grant = authorize(
        db,
        admin,
        "OFFICIAL_DORMITORY_MEMBER_UPDATE",
        [dormitory["management_grade_id"], new_grade],
    )
    if dormitory["version"] != version:
        raise ApiError(409, "DORMITORY_CONTENT_CONFLICT", "宿舍信息已变化，请重新载入")
    before = [member["imported_login_identifier"] for member in member_rows(db, [dormitory_id])]
    timestamp = now()
    db.execute(text("DELETE FROM official_dormitory_members WHERE official_dormitory_id=:id"), {"id": dormitory_id})
    for position, member in enumerate(members, 1):
        db.execute(
            text(
                """INSERT INTO official_dormitory_members(
                official_dormitory_id,user_id,role,position,imported_login_identifier,
                name_snapshot,grade_snapshot,major_snapshot,created_at)
                VALUES(:dormitory,:user,:role,:position,:login,:name,:grade,:major,:now)"""
            ),
            {
                "dormitory": dormitory_id,
                "user": member["id"],
                "role": "LEADER" if position == 1 else "MEMBER",
                "position": position,
                "login": member["login_identifier"],
                "name": member["name"],
                "grade": member["grade"],
                "major": member["major"],
                "now": timestamp,
            },
        )
    db.execute(
        text(
            """UPDATE official_dormitories SET management_grade_id=:grade,
            version=version+1,updated_at=:now WHERE id=:id"""
        ),
        {"grade": new_grade, "now": timestamp, "id": dormitory_id},
    )
    audit(
        db,
        admin,
        request,
        "UPDATE_OFFICIAL_DORMITORY_MEMBERS",
        "OFFICIAL_DORMITORY",
        dormitory_id,
        reason,
        grant=grant,
        before={"members": before},
        after={"members": [member["login_identifier"] for member in members]},
    )
    db.commit()
    return {"ok": True, "version": dormitory["version"] + 1}


@admin_router.post("/{dormitory_id}/content-moderation")
def moderate_official_dormitory_content(dormitory_id: int, request: Request, body: dict, db: DB) -> dict:
    admin = management_user(request, db)
    enforce_write_rate(request, admin["id"], "moderation", 20, 100, 60)
    field = body.get("field")
    action = body.get("action")
    if field not in CONTENT_FIELDS or action not in ("HIDE", "RESTORE", "RESET"):
        raise ApiError(400, "INVALID_DORMITORY_MODERATION", "治理操作无效")
    reason = clean_text(body.get("reason"), 200, True)
    begin_immediate(db)
    dormitory, grant = admin_dormitory(db, admin, dormitory_id, "OFFICIAL_DORMITORY_MODERATE")
    revision_field, column, status_column, _ = CONTENT_FIELDS[field]
    before = {"status": dormitory[status_column], "version": dormitory["version"]}
    timestamp = now()
    if action == "RESET":
        if dormitory[column]:
            db.execute(
                text(
                    """INSERT INTO official_dormitory_revisions(
                    official_dormitory_id,field_name,previous_value,new_value,from_version,to_version,
                    edited_by,editor_name_snapshot,created_at)
                    VALUES(:dormitory,:field,:previous,'',:version,:next_version,:editor,:name,:now)"""
                ),
                {
                    "dormitory": dormitory_id,
                    "field": revision_field,
                    "previous": dormitory[column],
                    "version": dormitory["version"],
                    "next_version": dormitory["version"] + 1,
                    "editor": admin["id"],
                    "name": admin["name"],
                    "now": timestamp,
                },
            )
        db.execute(
            text(
                f"UPDATE official_dormitories SET {column}='',{status_column}='NORMAL',version=version+1,updated_at=:now WHERE id=:id"
            ),
            {"now": timestamp, "id": dormitory_id},
        )
    else:
        status = "HIDDEN" if action == "HIDE" else "NORMAL"
        db.execute(
            text(
                f"UPDATE official_dormitories SET {status_column}=:status,version=version+1,updated_at=:now WHERE id=:id"
            ),
            {"status": status, "now": timestamp, "id": dormitory_id},
        )
    audit(
        db,
        admin,
        request,
        f"{action}_OFFICIAL_DORMITORY_CONTENT",
        "OFFICIAL_DORMITORY",
        dormitory_id,
        reason,
        metadata={"field": revision_field},
        grant=grant,
        before=before,
        after={"status": "NORMAL" if action in ("RESTORE", "RESET") else "HIDDEN"},
    )
    db.commit()
    return {"ok": True}


@admin_router.delete("/{dormitory_id}")
def delete_official_dormitory(dormitory_id: int, request: Request, body: dict, db: DB) -> dict:
    admin = management_user(request, db)
    grant = require_super_admin(admin)
    begin_immediate(db)
    dormitory = dormitory_by_id(db, dormitory_id)
    if not dormitory:
        raise ApiError(404, "OFFICIAL_DORMITORY_NOT_FOUND", ADMIN_NOT_FOUND_MESSAGE)
    if body.get("confirmation") != dormitory["dormitory_code"]:
        raise ApiError(400, "CONFIRMATION_REQUIRED", "请输入宿舍编号确认删除")
    reason = clean_text(body.get("reason"), 200, True)
    audit(
        db,
        admin,
        request,
        "DELETE_OFFICIAL_DORMITORY",
        "OFFICIAL_DORMITORY",
        dormitory_id,
        reason,
        grant=grant,
        before={"dormitoryCode": dormitory["dormitory_code"]},
    )
    db.execute(text("DELETE FROM official_dormitories WHERE id=:id"), {"id": dormitory_id})
    db.commit()
    return {"ok": True}
